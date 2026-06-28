import csv
import io
import json
import os
from datetime import datetime

from django.apps import apps
from django.conf import settings
from django.core import serializers
from django.db import transaction

from apps.core.exceptions import ConflictError, NotFoundError

from .models import Backup


class BackupService:

    BACKUP_DIR = settings.BASE_DIR / "backups"

    @classmethod
    def _ensure_backup_dir(cls):
        os.makedirs(cls.BACKUP_DIR, exist_ok=True)

    @classmethod
    def create_backup(cls, *, user, notes=""):
        cls._ensure_backup_dir()

        backup = Backup.objects.create(
            created_by=user,
            status=Backup.Status.IN_PROGRESS,
            notes=notes,
        )

        try:
            models_to_backup = [
                "properties.Property",
                "sections.Section",
                "units.Unit",
                "units.PricingRule",
                "occupants.Student",
                "occupancy.Occupancy",
                "payments.Payment",
                "payments.Receipt",
            ]

            data = {}
            record_counts = {}

            for model_label in models_to_backup:
                model = apps.get_model(model_label)
                qs = model.objects.all()
                record_counts[model_label] = qs.count()
                data[model_label] = json.loads(serializers.serialize("json", qs))

            metadata = {
                "record_counts": record_counts,
                "total_records": sum(record_counts.values()),
                "created_at": datetime.utcnow().isoformat(),
                "django_version": settings.DJANGO_VERSION if hasattr(settings, "DJANGO_VERSION") else "",
            }

            filename = f"backup_{backup.id:04d}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.json"
            file_path = cls.BACKUP_DIR / filename

            with open(file_path, "w") as f:
                json.dump({
                    "backup_id": backup.id,
                    "metadata": metadata,
                    "data": data,
                }, f, indent=2)

            file_size = os.path.getsize(file_path)

            backup.status = Backup.Status.COMPLETED
            backup.file_path = str(file_path)
            backup.file_size = file_size
            backup.metadata = metadata
            backup.save(update_fields=["status", "file_path", "file_size", "metadata"])

            return backup

        except Exception as e:
            backup.status = Backup.Status.FAILED
            backup.metadata = {"error": str(e)}
            backup.save(update_fields=["status", "metadata"])
            raise

    @staticmethod
    def list_backups(page=1, page_size=20):
        qs = Backup.objects.select_related("created_by").all()

        total = qs.count()
        start = (page - 1) * page_size
        end = start + page_size
        results = qs[start:end]

        return {
            "count": total,
            "page": page,
            "page_size": page_size,
            "total_pages": (total + page_size - 1) // page_size,
            "results": list(results),
        }

    @staticmethod
    def get_backup(pk):
        try:
            return Backup.objects.select_related("created_by").get(id=pk)
        except Backup.DoesNotExist:
            raise NotFoundError("Backup not found.")

    @staticmethod
    def validate_backup(backup):
        if backup.status != Backup.Status.COMPLETED:
            raise ConflictError(
                f"Cannot restore from a backup with status '{backup.status}'. "
                "Only completed backups can be restored."
            )
        if not backup.file_path or not os.path.exists(backup.file_path):
            raise ConflictError("Backup file not found on disk.")

        try:
            with open(backup.file_path) as f:
                data = json.load(f)
        except (json.JSONDecodeError, IOError) as e:
            raise ConflictError(f"Backup file is corrupted or unreadable: {e}")

        if "data" not in data:
            raise ConflictError("Invalid backup file: missing 'data' key.")

        return data

    @staticmethod
    @transaction.atomic
    def restore_backup(backup):
        data = BackupService.validate_backup(backup)

        models_to_restore = [
            "properties.Property",
            "sections.Section",
            "units.Unit",
            "units.PricingRule",
            "occupants.Student",
            "occupancy.Occupancy",
            "payments.Payment",
            "payments.Receipt",
        ]

        restored_counts = {}

        for model_label in models_to_restore:
            raw_data = data.get("data", {}).get(model_label, [])
            if not raw_data:
                restored_counts[model_label] = 0
                continue

            model = apps.get_model(model_label)

            for entry in raw_data:
                fields = entry.get("fields", {})
                pk = entry.get("pk")

                # Handle FK fields: convert to int if present
                for fk_field in ["student", "unit", "property", "section", "payment", "created_by", "actor"]:
                    if fk_field in fields and fields[fk_field] is not None:
                        fields[fk_field] = int(fields[fk_field])

                existing = model.objects.filter(pk=pk).first()
                if existing:
                    for field, value in fields.items():
                        setattr(existing, field, value)
                    existing.save()
                else:
                    model.objects.create(pk=pk, **fields)

            restored_counts[model_label] = len(raw_data)

        return restored_counts


class ExportService:

    @staticmethod
    def export_occupants(file_format="csv"):
        Student = apps.get_model("occupants.Student")
        qs = Student.objects.all().values(
            "id", "first_name", "last_name", "email", "phone",
            "student_id_number", "national_id", "is_active",
            "created_at", "updated_at",
        )
        return ExportService._write_export(qs, [
            "ID", "First Name", "Last Name", "Email", "Phone",
            "Student ID", "National ID", "Active",
            "Created At", "Updated At",
        ], file_format, "occupants")

    @staticmethod
    def export_occupancies(file_format="csv"):
        Occupancy = apps.get_model("occupancy.Occupancy")
        qs = Occupancy.objects.select_related("student", "unit__section__property").all()
        rows = []
        for o in qs:
            rows.append({
                "id": o.id,
                "student": str(o.student),
                "student_id": o.student_id,
                "unit": o.unit.name if o.unit else "",
                "unit_id": o.unit_id,
                "property": o.unit.section.property.name if o.unit and o.unit.section else "",
                "start_date": str(o.start_date),
                "end_date": str(o.end_date) if o.end_date else "",
                "billing_mode": o.billing_mode,
                "agreed_price": float(o.agreed_price) if o.agreed_price else "",
                "is_active": o.is_active,
                "created_at": str(o.created_at) if o.created_at else "",
            })
        return ExportService._write_export(rows, [
            "ID", "Student", "Student ID", "Unit", "Unit ID", "Property",
            "Start Date", "End Date", "Billing Mode", "Agreed Price",
            "Active", "Created At",
        ], file_format, "occupancies")

    @staticmethod
    def export_payments(file_format="csv"):
        Payment = apps.get_model("payments.Payment")
        qs = Payment.objects.select_related("student").all().values(
            "id", "student__first_name", "student__last_name",
            "student__student_id_number", "amount", "payment_date",
            "payment_method", "reference", "notes", "created_at", "updated_at",
        )
        rows = []
        for p in qs:
            rows.append({
                "id": p["id"],
                "student_name": f"{p['student__first_name']} {p['student__last_name']}".strip(),
                "student_id_number": p["student__student_id_number"] or "",
                "amount": float(p["amount"]),
                "payment_date": str(p["payment_date"]),
                "payment_method": p["payment_method"],
                "reference": p["reference"],
                "notes": p["notes"],
                "created_at": str(p["created_at"]) if p["created_at"] else "",
            })
        return ExportService._write_export(rows, [
            "ID", "Student Name", "Student ID Number", "Amount",
            "Payment Date", "Payment Method", "Reference", "Notes", "Created At",
        ], file_format, "payments")

    @staticmethod
    def export_receipts(file_format="csv"):
        Receipt = apps.get_model("payments.Receipt")
        qs = Receipt.objects.all().values(
            "id", "receipt_number", "issued_at",
            "outstanding_balance_after", "student_name",
            "student_id_number", "unit_name", "property_name",
            "amount", "payment_date", "payment_method", "reference",
        )
        rows = []
        for r in qs:
            rows.append({
                "id": r["id"],
                "receipt_number": r["receipt_number"],
                "issued_at": str(r["issued_at"]) if r["issued_at"] else "",
                "outstanding_balance_after": float(r["outstanding_balance_after"]),
                "student_name": r["student_name"],
                "student_id_number": r["student_id_number"] or "",
                "unit_name": r["unit_name"],
                "property_name": r["property_name"],
                "amount": float(r["amount"]),
                "payment_date": str(r["payment_date"]),
                "payment_method": r["payment_method"],
                "reference": r["reference"] or "",
            })
        return ExportService._write_export(rows, [
            "ID", "Receipt Number", "Issued At", "Outstanding Balance",
            "Student Name", "Student ID Number", "Unit", "Property",
            "Amount", "Payment Date", "Payment Method", "Reference",
        ], file_format, "receipts")

    @staticmethod
    def _write_export(rows, headers, file_format, entity_name):
        if file_format == "csv":
            return ExportService._write_csv(rows, headers, entity_name)
        elif file_format == "xlsx":
            return ExportService._write_xlsx(rows, headers, entity_name)
        raise ValueError(f"Unsupported format: {file_format}")

    @staticmethod
    def _write_csv(rows, headers, entity_name):
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([row.get(h.lower().replace(" ", "_").replace("-", "_")) for h in headers])
        content = output.getvalue()
        output.close()
        return content, f"{entity_name}.csv", "text/csv"

    @staticmethod
    def _sanitize_cell_value(value):
        if isinstance(value, datetime):
            return value.isoformat()
        return value

    @staticmethod
    def _write_xlsx(rows, headers, entity_name):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = entity_name

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill

        for row_idx, row in enumerate(rows, 2):
            for col_idx, header in enumerate(headers, 1):
                key = header.lower().replace(" ", "_").replace("-", "_")
                ws.cell(row=row_idx, column=col_idx, value=ExportService._sanitize_cell_value(row.get(key, "")))

        for col in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)

        output = io.BytesIO()
        wb.save(output)
        content = output.getvalue()
        output.close()
        return content, f"{entity_name}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
